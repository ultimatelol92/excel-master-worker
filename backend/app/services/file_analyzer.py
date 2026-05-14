from __future__ import annotations

import json
import logging
import os
import uuid
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

UPLOADS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")


class FileAnalyzer:
    def __init__(self) -> None:
        os.makedirs(UPLOADS_DIR, exist_ok=True)

    def save_upload(self, file_content: bytes, filename: str) -> tuple[str, str]:
        file_id = str(uuid.uuid4())
        safe_name = filename.replace(" ", "_")
        file_path = os.path.join(UPLOADS_DIR, f"{file_id}_{safe_name}")
        with open(file_path, "wb") as f:
            f.write(file_content)
        return file_id, file_path

    def analyze_file(self, file_path: str) -> dict[str, Any]:
        try:
            wb = load_workbook(file_path, data_only=True)
            result: dict[str, Any] = {
                "sheet_names": wb.sheetnames,
                "sheets": {},
            }

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_info = self._analyze_sheet(ws)

                wb_formulas = load_workbook(file_path)
                ws_formulas = wb_formulas[sheet_name]
                formulas = self._extract_formulas(ws_formulas, ws.max_row, ws.max_column)
                sheet_info["formulas"] = formulas
                wb_formulas.close()

                result["sheets"][sheet_name] = sheet_info

            wb.close()
            return result
        except Exception:
            logger.exception("Failed to analyze file: %s", file_path)
            raise

    @staticmethod
    def _analyze_sheet(ws: Any) -> dict[str, Any]:
        headers: list[str] = []
        for col in range(1, min(ws.max_column + 1, 53)):
            cell = ws.cell(row=1, column=col)
            if cell.value is not None:
                headers.append(str(cell.value))

        sample_data: list[list[str]] = []
        for row in range(2, min(ws.max_row + 1, 11)):
            row_data: list[str] = []
            for col in range(1, min(ws.max_column + 1, 53)):
                cell = ws.cell(row=row, column=col)
                row_data.append(str(cell.value) if cell.value is not None else "")
            sample_data.append(row_data)

        return {
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "headers": headers,
            "sample_data": sample_data,
        }

    @staticmethod
    def _extract_formulas(ws: Any, max_row: int, max_col: int) -> list[dict[str, str]]:
        formulas: list[dict[str, str]] = []
        for row in ws.iter_rows(min_row=1, max_row=min(max_row, 100), max_col=min(max_col, 52)):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                    })
        return formulas

    def get_file_summary_text(self, file_path: str) -> str:
        analysis = self.analyze_file(file_path)
        lines = [f"Excel File Analysis ({len(analysis['sheet_names'])} sheets):"]

        for sheet_name, info in analysis["sheets"].items():
            lines.append(f"\n--- Sheet: {sheet_name} ---")
            lines.append(f"Dimensions: {info['dimensions']} ({info['max_row']} rows x {info['max_column']} columns)")
            if info["headers"]:
                lines.append(f"Headers: {', '.join(info['headers'])}")
            if info["sample_data"]:
                lines.append("Sample data (first few rows):")
                for i, row_data in enumerate(info["sample_data"][:5], 2):
                    lines.append(f"  Row {i}: {json.dumps(row_data)}")
            if info["formulas"]:
                lines.append(f"Formulas found ({len(info['formulas'])}):")
                for f in info["formulas"][:20]:
                    lines.append(f"  {f['cell']}: {f['formula']}")

        return "\n".join(lines)
