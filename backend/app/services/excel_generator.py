from __future__ import annotations

import logging
import os
import uuid
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

logger = logging.getLogger(__name__)

GENERATED_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "generated")
UPLOADS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")


class ExcelGenerator:
    def __init__(self) -> None:
        os.makedirs(GENERATED_DIR, exist_ok=True)
        os.makedirs(UPLOADS_DIR, exist_ok=True)

    def generate_from_spec(self, spec: dict[str, Any]) -> tuple[str, str]:
        wb = Workbook()
        wb.remove(wb.active)

        sheets = spec.get("sheets", [])
        if not sheets:
            sheets = [{"name": "Sheet1", "columns": [], "data": []}]

        for sheet_spec in sheets:
            ws = wb.create_sheet(title=sheet_spec.get("name", "Sheet1"))
            self._build_sheet(ws, sheet_spec)

        named_ranges = spec.get("named_ranges", [])
        for nr in named_ranges:
            try:
                sheet_title = nr.get("sheet", sheets[0].get("name", "Sheet1"))
                cell_range = nr["range"]
                wb.defined_names.new(nr["name"], f"'{sheet_title}'!{cell_range}")
            except Exception:
                logger.warning("Failed to create named range: %s", nr.get("name"))

        file_name = spec.get("file_name", "generated.xlsx")
        if not file_name.endswith((".xlsx", ".xlsm")):
            file_name += ".xlsx"

        vba_macros = spec.get("vba_macros", [])
        if vba_macros:
            file_name = file_name.replace(".xlsx", ".xlsm")

        file_id = str(uuid.uuid4())
        file_path = os.path.join(GENERATED_DIR, f"{file_id}_{file_name}")

        if vba_macros:
            macro_info = self._format_macro_info(vba_macros)
            macro_ws = wb.create_sheet(title="VBA_Macros_Info")
            macro_ws["A1"] = "VBA Macros (copy to VBA Editor via Alt+F11)"
            macro_ws["A1"].font = Font(bold=True, size=14)
            row = 3
            for macro in vba_macros:
                macro_ws[f"A{row}"] = f"Macro: {macro.get('name', 'Unnamed')}"
                macro_ws[f"A{row}"].font = Font(bold=True, size=12, color="4472C4")
                row += 1
                for line in macro.get("code", "").split("\n"):
                    macro_ws[f"A{row}"] = line
                    macro_ws[f"A{row}"].font = Font(name="Consolas", size=10)
                    row += 1
                row += 1
            macro_ws.column_dimensions["A"].width = 80
            logger.info("VBA macros included as reference sheet: %s", macro_info)

        wb.save(file_path)
        logger.info("Generated Excel file: %s", file_path)
        return file_id, file_name

    def modify_file(self, file_id: str, modifications: dict[str, Any]) -> tuple[str, str]:
        source_path = self._find_file(file_id, UPLOADS_DIR)
        if not source_path:
            raise FileNotFoundError(f"Uploaded file not found: {file_id}")

        wb = load_workbook(source_path)

        for mod in modifications.get("modifications", []):
            sheet_name = mod.get("sheet")
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.active

            for formula in mod.get("add_formulas", []):
                try:
                    ws[formula["cell"]] = formula["formula"]
                except Exception:
                    logger.warning("Failed to add formula at %s", formula.get("cell"))

            for col_spec in mod.get("add_columns", []):
                col_idx = ws.max_column + 1
                col_letter = get_column_letter(col_idx)
                ws[f"{col_letter}1"] = col_spec.get("header", "")
                ws[f"{col_letter}1"].font = Font(bold=True)
                if col_spec.get("width"):
                    ws.column_dimensions[col_letter].width = col_spec["width"]

            for chart_spec in mod.get("add_charts", []):
                self._add_chart(ws, chart_spec)

            for cf_spec in mod.get("add_conditional_formatting", []):
                self._add_conditional_format(ws, cf_spec)

            for dv_spec in mod.get("add_data_validation", []):
                self._add_data_validation(ws, dv_spec)

            formatting = mod.get("add_formatting", {})
            if formatting:
                self._apply_formatting(ws, formatting)

        vba_macros = modifications.get("vba_macros", [])
        if vba_macros:
            macro_ws = wb.create_sheet(title="VBA_Macros_Info")
            macro_ws["A1"] = "VBA Macros (copy to VBA Editor via Alt+F11)"
            macro_ws["A1"].font = Font(bold=True, size=14)
            row = 3
            for macro in vba_macros:
                macro_ws[f"A{row}"] = f"Macro: {macro.get('name', 'Unnamed')}"
                macro_ws[f"A{row}"].font = Font(bold=True, size=12, color="4472C4")
                row += 1
                for line in macro.get("code", "").split("\n"):
                    macro_ws[f"A{row}"] = line
                    macro_ws[f"A{row}"].font = Font(name="Consolas", size=10)
                    row += 1
                row += 1
            macro_ws.column_dimensions["A"].width = 80

        file_name = modifications.get("file_name", "modified.xlsx")
        new_file_id = str(uuid.uuid4())
        new_path = os.path.join(GENERATED_DIR, f"{new_file_id}_{file_name}")
        wb.save(new_path)
        logger.info("Modified Excel file saved: %s", new_path)
        return new_file_id, file_name

    def analyze_uploaded_file(self, file_path: str) -> dict[str, Any]:
        wb = load_workbook(file_path, data_only=True)
        summary: dict[str, Any] = {
            "sheet_names": wb.sheetnames,
            "sheets": {},
        }

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_info: dict[str, Any] = {
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "headers": [],
                "sample_data": [],
                "formulas": [],
            }

            for col in range(1, min(ws.max_column + 1, 27)):
                cell = ws.cell(row=1, column=col)
                if cell.value is not None:
                    sheet_info["headers"].append(str(cell.value))

            for row in range(2, min(ws.max_row + 1, 6)):
                row_data = []
                for col in range(1, min(ws.max_column + 1, 27)):
                    cell = ws.cell(row=row, column=col)
                    row_data.append(str(cell.value) if cell.value is not None else "")
                sheet_info["sample_data"].append(row_data)

            wb_formulas = load_workbook(file_path)
            ws_formulas = wb_formulas[sheet_name]
            for row in ws_formulas.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        sheet_info["formulas"].append({
                            "cell": cell.coordinate,
                            "formula": cell.value,
                        })
            wb_formulas.close()

            summary["sheets"][sheet_name] = sheet_info

        wb.close()
        return summary

    def _build_sheet(self, ws: Any, spec: dict[str, Any]) -> None:
        columns = spec.get("columns", [])
        formatting = spec.get("formatting", {})
        header_color = formatting.get("header_color", "4472C4")
        header_font_color = formatting.get("header_font_color", "FFFFFF")

        header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        header_font = Font(bold=True, color=header_font_color, size=11)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, col_spec in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            header = col_spec if isinstance(col_spec, str) else col_spec.get("header", f"Col {col_idx}")
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            width = col_spec.get("width", 15) if isinstance(col_spec, dict) else 15
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        data = spec.get("data", [])
        alt_color = formatting.get("alternate_row_color", "")
        alt_fill = PatternFill(start_color=alt_color, end_color=alt_color, fill_type="solid") if alt_color else None

        for row_idx, row_data in enumerate(data, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(value, str) and value.startswith("="):
                    cell.value = value
                elif isinstance(value, (int, float)):
                    cell.value = value
                else:
                    try:
                        cell.value = float(value) if value and "." in str(value) else int(value)
                    except (ValueError, TypeError):
                        cell.value = value
                cell.border = thin_border
                if alt_fill and row_idx % 2 == 0:
                    cell.fill = alt_fill

        for formula_spec in spec.get("formulas", []):
            try:
                ws[formula_spec["cell"]] = formula_spec["formula"]
                ws[formula_spec["cell"]].border = thin_border
            except Exception:
                logger.warning("Failed to set formula at %s", formula_spec.get("cell"))

        number_formats = formatting.get("number_formats", {})
        for cell_range, fmt in number_formats.items():
            try:
                for row in ws[cell_range]:
                    if isinstance(row, tuple):
                        for cell in row:
                            cell.number_format = fmt
                    else:
                        row.number_format = fmt
            except Exception:
                logger.warning("Failed to apply number format to %s", cell_range)

        for merge_range in spec.get("merge_cells", []):
            try:
                ws.merge_cells(merge_range)
            except Exception:
                logger.warning("Failed to merge cells: %s", merge_range)

        freeze = spec.get("freeze_panes")
        if freeze:
            ws.freeze_panes = freeze

        for chart_spec in spec.get("charts", []):
            self._add_chart(ws, chart_spec)

        for cf_spec in spec.get("conditional_formatting", []):
            self._add_conditional_format(ws, cf_spec)

        for dv_spec in spec.get("data_validation", []):
            self._add_data_validation(ws, dv_spec)

    def _add_chart(self, ws: Any, chart_spec: dict[str, Any]) -> None:
        try:
            chart_type = chart_spec.get("type", "bar").lower()
            chart_map = {"bar": BarChart, "line": LineChart, "pie": PieChart}
            chart_cls = chart_map.get(chart_type, BarChart)
            chart = chart_cls()

            chart.title = chart_spec.get("title", "Chart")
            if hasattr(chart, "x_axis"):
                chart.x_axis.title = chart_spec.get("x_axis", "")
            if hasattr(chart, "y_axis"):
                chart.y_axis.title = chart_spec.get("y_axis", "")
            chart.width = chart_spec.get("width", 15)
            chart.height = chart_spec.get("height", 10)

            data_range = chart_spec.get("data_range", "A1:B10")
            parts = data_range.replace(":", "").split("$") if "$" in data_range else None

            if parts is None:
                col_start = ord(data_range[0].upper()) - ord("A") + 1
                sep_idx = data_range.index(":")
                col_end = ord(data_range[sep_idx + 1].upper()) - ord("A") + 1
                row_start = int("".join(c for c in data_range[1:sep_idx] if c.isdigit()))
                row_end = int("".join(c for c in data_range[sep_idx + 2:] if c.isdigit()))

                data_ref = Reference(ws, min_col=col_start + 1, min_row=row_start, max_col=col_end, max_row=row_end)
                cats = Reference(ws, min_col=col_start, min_row=row_start + 1, max_row=row_end)

                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats)

            position = chart_spec.get("position", "E2")
            ws.add_chart(chart, position)
        except Exception:
            logger.exception("Failed to add chart")

    @staticmethod
    def _add_conditional_format(ws: Any, cf_spec: dict[str, Any]) -> None:
        try:
            cell_range = cf_spec.get("range", "A1:A100")
            criteria = cf_spec.get("criteria", ">")
            value = cf_spec.get("value", 0)
            fmt_spec = cf_spec.get("format", {})

            fill = PatternFill(
                start_color=fmt_spec.get("bg_color", "92D050"),
                end_color=fmt_spec.get("bg_color", "92D050"),
                fill_type="solid",
            )
            font = Font(
                color=fmt_spec.get("font_color", "000000"),
                bold=fmt_spec.get("bold", False),
            )

            operator_map = {
                ">": "greaterThan",
                "<": "lessThan",
                ">=": "greaterThanOrEqual",
                "<=": "lessThanOrEqual",
                "=": "equal",
                "==": "equal",
                "!=": "notEqual",
                "between": "between",
            }
            operator = operator_map.get(criteria, "greaterThan")

            rule = CellIsRule(operator=operator, formula=[str(value)], fill=fill, font=font)
            ws.conditional_formatting.add(cell_range, rule)
        except Exception:
            logger.warning("Failed to add conditional formatting to %s", cf_spec.get("range"))

    @staticmethod
    def _add_data_validation(ws: Any, dv_spec: dict[str, Any]) -> None:
        try:
            cell_range = dv_spec.get("range", "A1:A100")
            dv_type = dv_spec.get("type", "list")

            if dv_type == "list":
                values = dv_spec.get("values", [])
                formula = ",".join(str(v) for v in values)
                dv = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
                dv.error = "Please select a valid option"
                dv.errorTitle = "Invalid Input"
                ws.add_data_validation(dv)
                dv.add(cell_range)
            elif dv_type in ("whole", "decimal"):
                dv = DataValidation(
                    type=dv_type,
                    operator=dv_spec.get("operator", "between"),
                    formula1=str(dv_spec.get("min", 0)),
                    formula2=str(dv_spec.get("max", 100)),
                )
                ws.add_data_validation(dv)
                dv.add(cell_range)
        except Exception:
            logger.warning("Failed to add data validation to %s", dv_spec.get("range"))

    @staticmethod
    def _apply_formatting(ws: Any, formatting: dict[str, Any]) -> None:
        number_formats = formatting.get("number_formats", {})
        for cell_range, fmt in number_formats.items():
            try:
                for row in ws[cell_range]:
                    if isinstance(row, tuple):
                        for cell in row:
                            cell.number_format = fmt
                    else:
                        row.number_format = fmt
            except Exception:
                logger.warning("Failed to apply number format to %s", cell_range)

    @staticmethod
    def _find_file(file_id: str, directory: str) -> str | None:
        for fname in os.listdir(directory):
            if fname.startswith(file_id):
                return os.path.join(directory, fname)
        return None

    @staticmethod
    def _format_macro_info(macros: list[dict[str, Any]]) -> str:
        return ", ".join(m.get("name", "Unnamed") for m in macros)

    @classmethod
    def get_generated_file_path(cls, file_id: str) -> str | None:
        return cls._find_file(file_id, GENERATED_DIR)

    @classmethod
    def get_uploaded_file_path(cls, file_id: str) -> str | None:
        return cls._find_file(file_id, UPLOADS_DIR)
